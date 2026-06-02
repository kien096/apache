from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
SOURCE_MD = DOCS / "technical-report.md"
OUTPUT_DOCX = DOCS / "BT01_Phan_tich_log_web_quy_mo_lon.docx"

CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGINS_DXA = {"top": 80, "bottom": 80, "start": 120, "end": 120}

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
MUTED = "667085"
LIGHT_GRAY = "F2F4F7"
CALLOUT_FILL = "F4F6F9"
PLACEHOLDER_FILL = "FFF4CC"
BORDER = "D0D5DD"


def set_page(section):
    section.start_type = WD_SECTION_START.NEW_PAGE
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)


def set_spacing(paragraph, before=0, after=6, line=1.10):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    fmt.line_spacing = line


def set_run_font(run, size=11, color="000000", bold=False, italic=False, font="Calibri"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Trang ")
    set_run_font(run, size=9, color=MUTED)
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_1)
    run._r.append(instr_text)
    run._r.append(fld_char_2)


def style_document(doc: Document):
    set_page(doc.sections[0])

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string("101828")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        style.paragraph_format.line_spacing = 1.10

    code = styles.add_style("Code Block", 1)
    code.font.name = "Consolas"
    code._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    code.font.size = Pt(9)
    code.font.color.rgb = RGBColor.from_string("111827")
    code.paragraph_format.space_before = Pt(2)
    code.paragraph_format.space_after = Pt(2)
    code.paragraph_format.line_spacing = 1.0

    caption = styles.add_style("BT01 Caption", 1)
    caption.font.name = "Calibri"
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor.from_string(MUTED)
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(4)

    for section in doc.sections:
        header = section.header.paragraphs[0]
        header.text = "BT01 - Phân tích log web quy mô lớn"
        set_spacing(header, after=0, line=1.0)
        header.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_run_font(header.runs[0], size=9, color=MUTED)
        footer = section.footer.paragraphs[0]
        add_page_number(footer)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, margins=CELL_MARGINS_DXA):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in margins.items():
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_grid(table, widths_dxa):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tbl.append(grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths_dxa[min(idx, len(widths_dxa) - 1)])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_borders(table, color=BORDER):
    borders = table._tbl.tblPr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        table._tbl.tblPr.append(borders)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def add_inline(paragraph, text, size=11, color="101828"):
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, size=10, color="111827", font="Consolas")
        elif part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, size=size, color=color, bold=True)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, size=size, color=color)


def add_table(doc, rows):
    if not rows:
        return
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [CONTENT_WIDTH_DXA // cols] * cols
    widths[-1] += CONTENT_WIDTH_DXA - sum(widths)
    set_table_grid(table, widths)
    set_table_borders(table)
    for r_idx, row in enumerate(rows):
        for c_idx in range(cols):
            cell = table.cell(r_idx, c_idx)
            text = row[c_idx] if c_idx < len(row) else ""
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            add_inline(p, text.strip(), size=10)
            if r_idx == 0:
                set_cell_shading(cell, LIGHT_GRAY)
                for run in p.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string(INK)
    doc.add_paragraph("")


def add_code_block(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_grid(table, [CONTENT_WIDTH_DXA])
    set_table_borders(table, color="E4E7EC")
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT_GRAY)
    set_cell_margins(cell, {"top": 100, "bottom": 100, "start": 140, "end": 140})
    cell.text = ""
    for idx, line in enumerate(lines):
        p = cell.paragraphs[0] if idx == 0 else cell.add_paragraph()
        set_spacing(p, before=0, after=0, line=1.0)
        run = p.add_run(line)
        set_run_font(run, size=9, color="111827", font="Consolas")
    doc.add_paragraph("")


def add_callout(doc, text, fill=CALLOUT_FILL, title=None):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_grid(table, [CONTENT_WIDTH_DXA])
    set_table_borders(table, color="E4E7EC")
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, {"top": 120, "bottom": 120, "start": 160, "end": 160})
    cell.text = ""
    p = cell.paragraphs[0]
    set_spacing(p, after=2, line=1.10)
    if title:
        add_inline(p, title, size=10, color=INK)
        p.runs[0].bold = True
        p = cell.add_paragraph()
        set_spacing(p, after=0, line=1.10)
    add_inline(p, text, size=10, color="344054")
    doc.add_paragraph("")


def add_image_placeholder(doc, title, description):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_grid(table, [CONTENT_WIDTH_DXA])
    set_table_borders(table, color="D6B656")
    cell = table.cell(0, 0)
    set_cell_shading(cell, PLACEHOLDER_FILL)
    set_cell_margins(cell, {"top": 260, "bottom": 260, "start": 180, "end": 180})
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_spacing(p, after=4, line=1.10)
    add_inline(p, title, size=11, color="7A5A00")
    for run in p.runs:
        run.bold = True
    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_spacing(p2, after=0, line=1.10)
    add_inline(p2, description, size=10, color="7A5A00")
    cap = doc.add_paragraph(style="BT01 Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.add_run(title)
    doc.add_paragraph("")


def add_cover(doc):
    for _ in range(4):
        doc.add_paragraph("")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("BÁO CÁO KỸ THUẬT")
    set_run_font(run, size=18, color=INK, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("BT01 - PHÂN TÍCH LOG WEB QUY MÔ LỚN")
    set_run_font(run, size=24, color=BLUE, bold=True)
    set_spacing(p, before=8, after=14, line=1.10)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("HDFS · MapReduce · YARN · Python Streaming · PostgreSQL · Grafana")
    set_run_font(run, size=12, color=MUTED)
    set_spacing(p, after=24, line=1.10)

    table = doc.add_table(rows=5, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_grid(table, [2600, CONTENT_WIDTH_DXA - 2600])
    set_table_borders(table)
    rows = [
        ("Lớp", "........................................................"),
        ("Môn học", "........................................................"),
        ("Giảng viên", "........................................................"),
        ("Thành viên nhóm", "Lại Thị Hồng Gấm - Lê Bá Kiên"),
        ("Ngày nộp", "........................................................"),
    ]
    for r_idx, (left, right) in enumerate(rows):
        for c_idx, value in enumerate([left, right]):
            cell = table.cell(r_idx, c_idx)
            set_cell_margins(cell)
            if c_idx == 0:
                set_cell_shading(cell, LIGHT_GRAY)
            cell.text = ""
            para = cell.paragraphs[0]
            add_inline(para, value, size=11, color=INK if c_idx == 0 else "101828")
            if c_idx == 0:
                para.runs[0].bold = True

    for _ in range(6):
        doc.add_paragraph("")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Ghi chú: Các khung màu vàng trong báo cáo là vị trí cần chèn ảnh minh họa/screenshot sau khi demo hệ thống.")
    set_run_font(run, size=10, color=MUTED, italic=True)
    doc.add_page_break()


def extract_headings(lines):
    headings = []
    for line in lines:
        if line.startswith("## "):
            headings.append((1, line[3:].strip()))
        elif line.startswith("### "):
            headings.append((2, line[4:].strip()))
    return headings


def add_static_toc(doc, headings):
    doc.add_heading("Mục Lục Tóm Tắt", level=1)
    for level, text in headings:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25 if level == 2 else 0)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(text)
        set_run_font(run, size=10 if level == 2 else 11, color=DARK_BLUE if level == 2 else INK, bold=(level == 1))
    doc.add_page_break()


def render_markdown(doc, text):
    lines = text.splitlines()
    add_static_toc(doc, extract_headings(lines))

    in_code = False
    code_lines = []
    table_lines = []
    quote_lines = []

    def flush_table():
        nonlocal table_lines
        if not table_lines:
            return
        rows = []
        for line in table_lines:
            if re.match(r"^\|\s*-+", line):
                continue
            row = [part.strip() for part in line.strip().strip("|").split("|")]
            rows.append(row)
        add_table(doc, rows)
        table_lines = []

    def flush_quote():
        nonlocal quote_lines
        if not quote_lines:
            return
        cleaned = [line.strip() for line in quote_lines if line.strip()]
        full = " ".join(cleaned)
        image_title = next(
            (
                line
                for line in cleaned
                if line.startswith("**[") and line.endswith("]**") and " - " in line
            ),
            None,
        )
        if image_title:
            desc = " ".join(line for line in cleaned if line != image_title)
            title = re.sub(r"^\*\*|\*\*$", "", image_title).strip()
            add_image_placeholder(doc, title, desc or "Chèn screenshot vào vị trí này.")
        else:
            add_callout(doc, full, fill=CALLOUT_FILL)
        quote_lines = []

    def flush_code():
        nonlocal code_lines
        if code_lines:
            add_code_block(doc, code_lines)
            code_lines = []

    for raw in lines:
        line = raw.rstrip()
        if line.startswith("```"):
            if in_code:
                flush_code()
                in_code = False
            else:
                flush_table()
                flush_quote()
                in_code = True
            continue
        if in_code:
            code_lines.append(line)
            continue
        if line.startswith("|") and line.endswith("|"):
            flush_quote()
            table_lines.append(line)
            continue
        flush_table()
        if line.startswith(">"):
            quote_lines.append(line[1:].strip())
            continue
        flush_quote()

        if not line.strip():
            continue
        if line.startswith("# "):
            continue
        if line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
        elif re.match(r"^- ", line):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Inches(0.5)
            p.paragraph_format.first_line_indent = Inches(-0.25)
            set_spacing(p, after=8, line=1.167)
            add_inline(p, line[2:].strip())
        elif re.match(r"^\d+\. ", line):
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.left_indent = Inches(0.5)
            p.paragraph_format.first_line_indent = Inches(-0.25)
            set_spacing(p, after=8, line=1.167)
            add_inline(p, re.sub(r"^\d+\. ", "", line).strip())
        else:
            p = doc.add_paragraph()
            set_spacing(p)
            add_inline(p, line)

    flush_table()
    flush_quote()
    flush_code()


def add_appendix_summary(doc):
    doc.add_page_break()
    doc.add_heading("Phiếu Kiểm Tra Yêu Cầu BT01", level=1)
    rows = [
        ["Yêu cầu", "Phần đã thực hiện", "Minh chứng cần nêu khi bảo vệ"],
        ["Cài đặt cụm Hadoop", "Docker Compose với NameNode, 2 DataNode, ResourceManager, 2 NodeManager", "Ảnh HDFS UI/YARN UI và file docker-compose.yml"],
        ["Tổ chức log theo ngày", "/logs/yyyy/MM/dd/access.log", "Ảnh HDFS UI tại /logs/2026/06/"],
        ["Top 20 IP", "Job top_ips_global", "Output /bt01/output/top_ips_global"],
        ["Mã HTTP theo ngày", "Job status_by_day", "Bảng status_by_day trong PostgreSQL/Grafana"],
        ["Top 10 URL mỗi ngày", "Job top_urls_by_day", "Bảng top_urls_by_day"],
        ["Lưu lượng theo 24 giờ", "Job traffic_by_hour", "Biểu đồ time series trong Grafana"],
        ["Theo dõi YARN", "Tất cả job chạy qua Hadoop Streaming trên YARN", "Application FINISHED trên http://localhost:8089"],
        ["Benchmark reducer", "Chạy 1, 2, 4 reducer", "docs/benchmark-results.csv và biểu đồ so sánh"],
        ["Dashboard tương tác", "Grafana đọc dữ liệu từ PostgreSQL", "Dashboard BT01 - Web Log Analysis"],
    ]
    add_table(doc, rows)


def main():
    text = SOURCE_MD.read_text(encoding="utf-8")
    doc = Document()
    style_document(doc)
    add_cover(doc)
    render_markdown(doc, text)
    add_appendix_summary(doc)
    doc.save(OUTPUT_DOCX)
    print(str(OUTPUT_DOCX.name))


if __name__ == "__main__":
    main()
